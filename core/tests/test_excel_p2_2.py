"""P2-2 回归：整份报告的 Excel 中间产物（对标 data-analyzer 分析结果.xlsx）。

正向：含离群值的数据集 → 原始数据 sheet 红标 + 异常清单有记录（系列级+行级）。
反向：去掉离群值（干净数据）→ 异常清单应为空，证明红标/清单不是「永远有」，非假绿。
边界：无图表（降级报告）→ 不崩溃，原始数据 sheet 仍在，透视 sheet 给提示。
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from openpyxl import load_workbook

from core.model import Column, ChartSpec, Dataset, ReportSpec
from core.excel_export import save_report_excel, RED_FILL_RGB


def _sales_ds(rid, name, sales):
    cols = [Column("地区", "text"), Column("销售额", "number")]
    rows = [{"地区": r, "销售额": v} for r, v in zip(
        ["华东", "华南", "华北", "西南", "西北"], sales)]
    return Dataset(id=rid, name=name, source_type="csv", columns=cols,
                   rows=rows, total_rows=len(rows))


def _spec():
    return ReportSpec(title="测试报告", charts=[
        ChartSpec(type="bar", x="地区", y=["销售额"], agg="sum")])


def _has_red_fill(ws):
    for row in ws.iter_rows():
        for cell in row:
            if RED_FILL_RGB and RED_FILL_RGB in str(getattr(cell.fill.fgColor, "rgb", "")):
                return True
    return False


def test_positive_with_outlier():
    ds = _sales_ds("a1", "销售A", [5000, 200, 150, 180, 220])  # 华东离群
    sub = {"datasets": [ds], "spec": _spec()}
    with tempfile.TemporaryDirectory() as d:
        p = save_report_excel(sub, d)
        assert os.path.exists(p), "正向：应生成 分析结果.xlsx"
        wb = load_workbook(p)
        sheets = wb.sheetnames
        assert any(s.startswith("原始数据_") for s in sheets), f"正向：应有原始数据 sheet，实际 {sheets}"
        assert "透视" in sheets, f"正向：应有透视 sheet，实际 {sheets}"
        assert "异常清单" in sheets, f"正向：应有异常清单 sheet，实际 {sheets}"
        # 红标：原始数据 sheet 至少一处红底
        raw = wb[next(s for s in sheets if s.startswith("原始数据_"))]
        assert _has_red_fill(raw), "正向：原始数据应出现红标单元格"
        # 异常清单：华东既在系列级又在行级出现
        anom = wb["异常清单"]
        rows = list(anom.iter_rows(min_row=2, values_only=True))
        rows = [r for r in rows if any(c is not None for c in r)]
        assert len(rows) >= 2, f"正向：异常清单应≥2条（系列+行级），实际 {len(rows)}：{rows}"
        assert any("华东" in str(r) for r in rows), "正向：异常清单应含离群的华东"
        # 透视 sheet 有聚合表（含「华东」与数值）
        pivot = wb["透视"]
        pivot_text = " ".join(str(c.value) for row in pivot.iter_rows() for c in row if c.value is not None)
        assert "华东" in pivot_text and "5000" in pivot_text, "正向：透视应含华东 5000"
    print("  正向通过：红标+异常清单+透视 均生成，异常条数", len(rows))


def test_negative_clean_data():
    """反向验证：干净数据（无离群）→ 异常清单必须为空，证明检测点真生效。"""
    ds = _sales_ds("b1", "销售B", [200, 210, 190, 205, 195])  # 全部落在 IQR 内
    sub = {"datasets": [ds], "spec": _spec()}
    with tempfile.TemporaryDirectory() as d:
        p = save_report_excel(sub, d)
        wb = load_workbook(p)
        anom = wb["异常清单"]
        rows = [r for r in anom.iter_rows(min_row=2, values_only=True)
                if any(c is not None for c in r)]
        assert len(rows) == 0, f"反向：干净数据异常清单应为空，实际 {len(rows)}：{rows}"
        raw = wb[next(s for s in wb.sheetnames if s.startswith("原始数据_"))]
        assert not _has_red_fill(raw), "反向：干净数据不应有任何红标"
    print("  反向验证通过（检测点确实生效，非假绿）")


def test_boundary_no_charts():
    """边界：无图表（降级报告）不应崩溃，原始数据仍在，透视给提示。"""
    ds = _sales_ds("c1", "销售C", [200, 210, 190, 205, 195])
    spec = ReportSpec(title="降级报告", charts=[])  # 没有图
    sub = {"datasets": [ds], "spec": spec}
    with tempfile.TemporaryDirectory() as d:
        p = save_report_excel(sub, d)
        wb = load_workbook(p)
        assert any(s.startswith("原始数据_") for s in wb.sheetnames), "边界：原始数据 sheet 应仍在"
        pivot = wb["透视"]
        pivot_text = " ".join(str(c.value) for row in pivot.iter_rows() for c in row if c.value is not None)
        assert "没有可聚合" in pivot_text, "边界：透视应提示无图表"
    print("  无图表边界通过（不崩溃）")


if __name__ == "__main__":
    test_positive_with_outlier()
    test_negative_clean_data()
    test_boundary_no_charts()
    print("ALL P2-2 TESTS PASSED")
