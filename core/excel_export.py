"""整份报告的 Excel 中间产物（对标 data-analyzer 的 分析结果.xlsx）。

三个维度，与 data-analyzer 的输出物对齐：
1. 原始数据_<数据集>：原始行 + ⚠异常标记列；行级离群单元格红标（detect_row_anomalies）。
2. 透视：每个图聚合后的「类别 × 指标」表（来自 aggregate）。
3. 异常清单：跨所有图与数据集的异常汇总（系列级 + 行级），方便逐条核查。

命名沿用 分析结果.xlsx，让用过 data-analyzer 的用户一眼认得这份产物。
"""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .anomaly import detect_row_anomalies
from .chartset import build_blocks

# 红标：浅红底 + 深红字，和 data-analyzer 的红高亮观感一致
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT = Font(color="9C0006", bold=True)
HEAD_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEAD_FONT = Font(bold=True)
RED_FILL_RGB = "00FFC7CE"  # openpyxl 读回时 fgColor.rgb 的写法（带 alpha 前缀）


def _safe_sheet(title: str, used: set) -> str:
    """ sheet 名 ≤31 字符且唯一。"""
    t = title[:31]
    if t not in used:
        used.add(t)
        return t
    i = 1
    while f"{t[:27]}_{i}" in used:
        i += 1
    out = f"{t[:27]}_{i}"
    used.add(out)
    return out


def _iter_aggs(datasets, spec):
    """复用 build_blocks 的口径，逐图产出 (数据集名, 图标题, agg)。
    blocks 与 pending 同序，用 block 的 dataset id 反查数据集名。
    """
    blocks, pending = build_blocks(datasets, spec.charts)
    id_to_name = {ds.id: ds.name for ds in datasets}
    for block, item in zip(blocks, pending):
        yield (id_to_name.get(block.get("dataset"), ""),
               item.get("spec", {}).title or "",
               item["agg"])


def _write_raw_sheet(wb, ds, used):
    ws = wb.create_sheet(_safe_sheet(f"原始数据_{ds.name}", used))
    headers = [c.name for c in ds.columns] + ["⚠异常标记"]
    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=ci)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
    row_anom = detect_row_anomalies(ds)
    name_idx = {c.name: i for i, c in enumerate(ds.columns)}
    for ri, r in enumerate(ds.rows):
        vals = [r.get(h) if h != "⚠异常标记" else "" for h in headers]
        ws.append(vals)
        markers = row_anom.get(ri)
        if markers:
            mark = "⚠ " + ";".join(col for col, _ in markers)
            cell = ws.cell(row=ri + 2, column=len(headers))
            cell.value = mark
            cell.font = RED_FONT
            for col, _ in markers:
                ws.cell(row=ri + 2, column=name_idx[col] + 1).fill = RED_FILL
    for ci, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(ci)].width = min(40, max(10, len(str(h)) + 2))
    return ws


def _write_pivot_sheet(wb, datasets, spec, used):
    ws = wb.create_sheet(_safe_sheet("透视", used))
    r = 1
    any_table = False
    for _ds_name, title, agg in _iter_aggs(datasets, spec):
        if agg.get("empty") or agg.get("scatter"):
            continue
        cats = agg.get("categories") or []
        if not cats:
            continue
        any_table = True
        ws.cell(row=r, column=1, value=title).font = HEAD_FONT
        r += 1
        series = agg.get("series", [])
        head = [agg.get("x") or "类别"] + [s["name"] for s in series]
        for ci, h in enumerate(head, 1):
            c = ws.cell(row=r, column=ci, value=h)
            c.fill = HEAD_FILL
            c.font = HEAD_FONT
        r += 1
        for j, cat in enumerate(cats):
            row = [cat] + [s["data"][j] if j < len(s["data"]) else "" for s in series]
            for ci, v in enumerate(row, 1):
                ws.cell(row=r, column=ci, value=v)
            r += 1
        r += 1  # 表间空行
    if not any_table:
        ws.cell(row=1, column=1, value="（本报告没有可聚合的图表）")


def _write_anomaly_sheet(wb, datasets, spec, used):
    ws = wb.create_sheet(_safe_sheet("异常清单", used))
    head = ["数据集", "来源", "图表/列", "位置", "数值", "类型", "说明"]
    ws.append(head)
    for ci in range(1, len(head) + 1):
        c = ws.cell(row=1, column=ci)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
    rows = []
    # 系列级：图表聚合后的类别点异常
    for ds_name, title, agg in _iter_aggs(datasets, spec):
        for s in agg.get("series", []):
            for a in s.get("anomalies", []):
                rows.append([ds_name, "图表", title,
                             a.get("category", ""), a.get("value"),
                             "/".join(a.get("methods", [])), a.get("reason", "")])
    # 行级：原始数据里的离群单元格
    for ds in datasets:
        for ri, marks in detect_row_anomalies(ds).items():
            for col, reason in marks:
                rows.append([ds.name, "原始数据", col, f"第{ri + 1}行",
                             ds.rows[ri].get(col), "IQR离群", reason])
    for row in rows:
        ws.append(row)
    return len(rows)


def save_report_excel(sub: dict, out_dir: str) -> str:
    """把一份已生成的报告导出成 分析结果.xlsx，返回文件路径。

    sub 需含 "datasets" 与 "spec"（与 /api/export 持有的 REPORTS 子项一致）。
    """
    datasets = sub["datasets"]
    spec = sub["spec"]
    wb = Workbook()
    wb.remove(wb.active)  # 删掉默认空 sheet
    used = set()
    for ds in datasets:
        _write_raw_sheet(wb, ds, used)
    _write_pivot_sheet(wb, datasets, spec, used)
    _write_anomaly_sheet(wb, datasets, spec, used)
    path = os.path.join(out_dir, "分析结果.xlsx")
    wb.save(path)
    return path
