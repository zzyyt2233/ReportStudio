"""文档格式互转：Word <-> PDF（纯 Python 实现，不依赖 LibreOffice / Office）。

能力边界（务必知悉，不是 bug）：
- docx -> pdf：python-docx 读 + reportlab 重建。保留段落 / 标题层级 / 加粗 /
  字体色 / 表格 / 内嵌图片；不保留分栏、文本框，页眉页脚水印，批注，SmartArt，
  公式(OME)。复杂排版会近似还原。
- pdf -> docx：pdfplumber 提取文本与表格 + python-docx 重建。仅支持**有文字层**的
  PDF（扫描件 / 图片型 PDF 无文字层，需 OCR，不在范围内）；多栏排版可能串行；
  图片不还原。表格区域会自动从正文里剔除，单独还原成表格。

这些都是纯 Python 重建格式的固有上限。转换是同步 CPU 活，路由层用同步 def 跑。
"""

from __future__ import annotations

import os
import tempfile
from xml.sax.saxutils import escape as _xml_escape

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
except ImportError:
    Document = None

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except ImportError:
    SimpleDocTemplate = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None


CN_FONT = "CN"

# 表格 / 网格用的项目灰，和前端 --line / 表头底保持一致
_GRID = colors.HexColor("#C9CCD1")
_HEAD = colors.HexColor("#EDF1F7")


def _register_font() -> bool:
    """注册中文字体：优先系统微软雅黑 / 黑体 / 宋体，都没有退回 reportlab 自带 CID。"""
    if "CN" in pdfmetrics.getRegisteredFontNames():
        return True
    windir = os.environ.get("WINDIR", r"C:\Windows")
    candidates = [
        (os.path.join(windir, "Fonts", "msyh.ttc"), 0),
        (os.path.join(windir, "Fonts", "msyhbd.ttc"), 0),
        (os.path.join(windir, "Fonts", "simhei.ttf"), 0),
        (os.path.join(windir, "Fonts", "simsun.ttc"), 0),
    ]
    for path, idx in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(CN_FONT, path, subfontIndex=idx))
                return True
            except Exception:
                continue
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        globals()["CN_FONT"] = "STSong-Light"
        return True
    except Exception:
        return False


# ------------------------------ docx -> pdf ------------------------------

def convert_docx_to_pdf(src: str, out: str) -> None:
    if Document is None or SimpleDocTemplate is None:
        raise RuntimeError("未安装 python-docx / reportlab，无法转换")
    _register_font()
    doc = Document(src)
    pdf = SimpleDocTemplate(
        out, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
        title=os.path.splitext(os.path.basename(src))[0],
    )
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName=CN_FONT,
                          fontSize=10.5, leading=16)
    h1 = ParagraphStyle("h1", parent=body, fontSize=17, leading=22,
                        spaceBefore=10, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=body, fontSize=14, leading=19,
                        spaceBefore=8, spaceAfter=4)
    h3 = ParagraphStyle("h3", parent=body, fontSize=12, leading=17,
                        spaceBefore=6, spaceAfter=3)
    title = ParagraphStyle("title", parent=h1, alignment=TA_CENTER)
    styles = {"Heading 1": h1, "Heading 2": h2, "Heading 3": h3, "Title": title}

    def _rgb(run):
        try:
            if run.font.color and run.font.color.type is not None:
                return f"#{run.font.color.rgb}"
        except Exception:
            return None
        return None

    def _run_text(run):
        t = _xml_escape(run.text or "")
        if not t:
            return ""
        color = _rgb(run)
        if color:
            t = f'<font color="{color}">{t}</font>'
        if run.bold:
            t = f"<b>{t}</b>"
        return t

    def _para_style(p):
        sn = (p.style.name if p.style else "") or ""
        return styles.get(sn, body)

    flow = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph as DocxParagraph
            p = DocxParagraph(child, doc)
            sn = (p.style.name if p.style else "") or ""
            text = "".join(_run_text(r) for r in p.runs)
            if sn.startswith("List"):
                text = "• " + text
            for im in _extract_images(p, doc):
                try:
                    flow.append(Image(im, width=16 * cm))
                    flow.append(Spacer(1, 4))
                except Exception:
                    pass
            if text.strip():
                flow.append(Paragraph(text, _para_style(p)))
        elif child.tag == qn("w:tbl"):
            from docx.table import Table as DocxTable
            t = DocxTable(child, doc)
            rows = [[_cell_text(c) for c in row.cells] for row in t.rows]
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            data = [r + [""] * (ncol - len(r)) for r in rows]
            tbl = Table(data)
            tbl.hAlign = "LEFT"
            tbl.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), CN_FONT),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.4, _GRID),
                ("BACKGROUND", (0, 0), (-1, 0), _HEAD),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]))
            flow.append(Spacer(1, 4))
            flow.append(tbl)
            flow.append(Spacer(1, 4))
    pdf.build(flow)


def _cell_text(cell) -> str:
    return (cell.text or "").replace("\n", " ").strip()


def _extract_images(p, doc) -> list[str]:
    """提取段落里内嵌图片，落临时文件，返回路径列表（顺序即文档顺序）。"""
    out = []
    for blip in p._p.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        part = doc.part.related_parts.get(rid)
        if not part:
            continue
        ct = getattr(part, "content_type", "") or ""
        ext = ".jpg" if "jpeg" in ct else ".png"
        fn = os.path.join(tempfile.gettempdir(), f"rs_cv_{abs(hash(rid))}{ext}")
        try:
            with open(fn, "wb") as fp:
                fp.write(part.blob)
            out.append(fn)
        except Exception:
            pass
    return out


# ------------------------------ pdf -> docx ------------------------------

def convert_pdf_to_docx(src: str, out: str) -> None:
    if pdfplumber is None or Document is None:
        raise RuntimeError("未安装 pdfplumber / python-docx，无法转换")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    with pdfplumber.open(src) as pdf:
        scanned = True
        for page in pdf.pages:
            tables = page.find_tables()
            tb = [t.bbox for t in tables]
            chars = page.chars
            if not chars:
                continue
            scanned = False
            # 按行聚类（top 四舍五入），保留横向顺序
            lines: dict[int, list] = {}
            for c in chars:
                lines.setdefault(round(c["top"]), []).append(c)
            for top in sorted(lines):
                row = lines[top]
                x0 = min(c["x0"] for c in row)
                up = min(c["top"] for c in row)
                bottom = max(c["bottom"] for c in row)
                if _in_any(tb, x0, up, bottom):
                    continue
                size = max(c.get("size", 10) for c in row)
                txt = "".join(c["text"] for c in sorted(row, key=lambda c: c["x0"])).strip()
                if not txt:
                    continue
                if size >= 14:
                    doc.add_heading(txt, level=1 if size >= 18 else 2)
                else:
                    doc.add_paragraph(txt)
            for t in tables:
                extracted = t.extract()
                if not extracted:
                    continue
                ncol = max(len(r) for r in extracted)
                tbl = doc.add_table(rows=len(extracted), cols=ncol)
                tbl.style = "Light Grid Accent 1"
                for i, r in enumerate(extracted):
                    for j in range(ncol):
                        v = r[j] if j < len(r) else None
                        tbl.rows[i].cells[j].text = "" if v is None else str(v)
        if scanned:
            raise RuntimeError("该 PDF 没有文字层（疑似扫描件 / 图片型），"
                               "纯文本提取无法还原内容，需要 OCR 支持")
    doc.save(out)


def _in_any(bboxes, x0, top, bottom) -> bool:
    """判断一行是否落在某个表格 bbox 内（用于从正文剔除表格区文字）。"""
    mid = (top + bottom) / 2
    for (x0b, topb, x1b, botb) in bboxes:
        if topb - 1 <= mid <= botb + 1 and x0b - 2 <= x0 <= x1b + 2:
            return True
    return False
